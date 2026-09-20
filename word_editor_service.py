from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def extract_docx(path):
    doc=Document(path); blocks=[]
    for p in doc.paragraphs:
        runs=[]
        for r in p.runs:runs.append({'text':r.text,'bold':bool(r.bold),'italic':bool(r.italic),'underline':bool(r.underline),'size':float(r.font.size.pt) if r.font.size else None,'color':str(r.font.color.rgb) if r.font.color and r.font.color.rgb else None})
        blocks.append({'type':'paragraph','text':p.text,'style':p.style.name if p.style else 'Normal','runs':runs})
    for t in doc.tables:blocks.append({'type':'table','rows':[[c.text for c in row.cells] for row in t.rows]})
    return {'blocks':blocks}
def save_docx(path,data):
    doc=Document()
    for b in data.get('blocks',[]):
        if b.get('type')=='table':
            rows=b.get('rows',[]);t=doc.add_table(rows=len(rows),cols=max([len(r) for r in rows] or [1]))
            for i,row in enumerate(rows):
                for j,v in enumerate(row):t.cell(i,j).text=str(v)
            continue
        p=doc.add_paragraph();
        try:p.style=b.get('style','Normal')
        except:pass
        p.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}.get(b.get('align','left'),WD_ALIGN_PARAGRAPH.LEFT)
        for d in b.get('runs') or [{'text':b.get('text','')}]:
            r=p.add_run(d.get('text',''));r.bold=d.get('bold');r.italic=d.get('italic');r.underline=d.get('underline')
            if d.get('size'):r.font.size=Pt(float(d['size']))
            if d.get('color'):
                try:r.font.color.rgb=RGBColor.from_string(str(d['color']).replace('#','')[-6:])
                except:pass
    doc.save(path)
